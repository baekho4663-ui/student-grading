import io
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from grading import ai_grade, demo_grade, export_excel, validate_result


class GradingTests(unittest.TestCase):
    def setUp(self):
        self.questions = [{"prompt": "문제", "rubric": "빛 2점", "max_score": 2, "required": ["빛"]}]

    def test_invalid_ai_scores_and_missing_elements_rejected(self):
        rows = demo_grade(self.questions)
        for score in [-1, 3, True, float("nan")]:
            rows[0]["score"] = score
            with self.assertRaises(ValueError):
                validate_result({"questions": rows}, self.questions)
        rows = demo_grade(self.questions)
        rows[0]["checks"] = []
        with self.assertRaises(ValueError):
            validate_result({"questions": rows}, self.questions)

    def test_excel_keeps_original_and_final_and_escapes_formula(self):
        initial = demo_grade(self.questions)
        final = demo_grade(self.questions)
        final[0]["score"] = 2
        final[0]["reason"] = "=1+1"
        data = export_excel("=HYPERLINK(1)", self.questions, initial, final, "관찰", "초안", "체험")
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ["평가 결과", "필수 요소", "세특 초안"])
        self.assertEqual(book["평가 결과"]["A2"].data_type, "s")
        self.assertEqual(book["평가 결과"]["E4"].value, 0)
        self.assertEqual(book["평가 결과"]["G4"].value, 2)
        self.assertEqual(book["평가 결과"]["H4"].value, "'=1+1")

    def test_ai_request_uses_image_and_validates_response(self):
        import json
        client = MagicMock()
        client.chat.completions.create.return_value.choices[0].message.content = json.dumps({"questions": demo_grade(self.questions)})
        with patch("grading.client_for", return_value=client):
            result = ai_grade("test-key", "test-model", self.questions, [{"bytes": b"image", "mime": "image/png"}])
        self.assertEqual(len(result), 1)
        kwargs = client.chat.completions.create.call_args.kwargs
        self.assertTrue(kwargs["messages"][1]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,"))
        self.assertEqual(kwargs["response_format"]["type"], "json_schema")

    def test_demo_teacher_confirmation_draft_and_invalidation(self):
        app = AppTest.from_file("app.py", default_timeout=20).run()
        self.assertFalse(app.exception)
        next(b for b in app.button if b.label == "체험 결과 만들기").click().run()
        self.assertFalse(app.exception)
        score = next(n for n in app.number_input if n.label == "교사 점수")
        score.set_value(6.0).run()
        next(c for c in app.checkbox if "모든 문항" in c.label).check().run()
        next(b for b in app.button if b.label == "교사 최종 확정").click().run()
        self.assertIn("confirmed", app.session_state)
        next(t for t in app.text_area if t.label.startswith("교사 관찰 기록")).set_value("실험 결과를 표로 정리함.").run()
        next(b for b in app.button if b.label == "세특 초안 만들기").click().run()
        self.assertIn("실험 결과", app.session_state["record"])
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get("download_button")), 1)
        next(t for t in app.text_area if t.label.startswith("교사 관찰 기록")).set_value("관찰 수정").run()
        self.assertNotIn("record", app.session_state)
        next(n for n in app.number_input if n.label == "교사 점수").set_value(5.0).run()
        self.assertNotIn("confirmed", app.session_state)
        self.assertEqual(len(app.get("download_button")), 0)
        next(t for t in app.text_area if t.label == "문제").set_value("새 문제").run()
        self.assertNotIn("initial", app.session_state)
        self.assertFalse(app.exception)


if __name__ == "__main__":
    unittest.main()
